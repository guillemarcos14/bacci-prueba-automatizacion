"""Motor determinista de importación, conciliación y priorización.

No modifica las fuentes ni interpreta un correo como confirmación del ERP.
"""

from __future__ import annotations

import hashlib
import json
import re
import sqlite3
import time
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ORDER_RE = re.compile(r"\bP-(?:X)?\d+\b", re.IGNORECASE)
LINE_RE = re.compile(r"l[ií]nea\s*(\d{4,6})\b", re.IGNORECASE)
SLASH_LINE_RE = re.compile(r"^\s*/\s*(\d{4,6})\b")
DATE_RE = re.compile(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b")
SKU_RE = re.compile(r"\b[A-Z]{2,4}-\d{2,4}\b", re.IGNORECASE)
SUPERSEDES_RE = re.compile(r"rectifico\s+mi\s+correo\s+(msg-\d+)", re.IGNORECASE)
INITIAL_CUTOFF = datetime(2026, 9, 10, 10, 0)
UPDATE_CUTOFF = datetime(2026, 9, 10, 11, 0)
REQUIRED_COLUMNS = {
    "Pedidos": {"pedido_id", "linea_id", "cliente_id", "sku", "color", "talla", "uds_pedidas",
                "uds_enviadas", "precio_unitario_eur", "fecha_compromiso"},
    "Clientes": {"cliente_id", "nombre", "email"},
    "Correos": {"message_id", "received_at", "from", "to", "subject", "body"},
}


def plain(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat(timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    return value


def canonical_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=plain)


def digest(value: Any) -> str:
    return hashlib.sha256(canonical_json(value).encode("utf-8")).hexdigest()


def fold(value: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", value.casefold()) if not unicodedata.combining(c))


def search_text(value: Any) -> str:
    """Indexa valores de origen, incluidos pedidos, maestro y correos, sin depender de nombres de columnas."""
    if isinstance(value, dict):
        return " ".join(search_text(item) for item in value.values())
    if isinstance(value, (list, tuple)):
        return " ".join(search_text(item) for item in value)
    return fold(str(value)) if value is not None else ""


def case_label(case: dict[str, Any]) -> str:
    """Estado visible del detalle; la prioridad solo ordena los casos activos."""
    if not case["attention"]:
        return "Servido"
    if case["unresolved"]:
        return "Sin correspondencia"
    if any(issue not in ("duplicado idéntico consolidado",) for issue in case["issues"]):
        return "Revisión humana"
    if case.get("request_count"):
        return "Solicitud sin confirmar"
    if case["priority"] == "Alta":
        return "Atención prioritaria"
    return "Pendiente"


def case_search_text(case: dict[str, Any], customer_email: str) -> str:
    # La prioridad interna de una línea servida no forma parte de su estado visible.
    internal = {"case_id", "kind", "priority", "priority_rank", "attention", "unresolved",
                "overdue", "due_today", "request_count", "request_conflict", "sort_time"}
    source = {key: value for key, value in case.items() if key not in internal}
    dates = [date.fromisoformat(value).strftime("%d/%m/%Y") for value in
             (case.get("fecha_compromiso"), case.get("requested_date")) if value]
    return search_text((source, customer_email, dates, case_label(case),
                        case["priority"] if case["attention"] else ""))


def load_sheet(path: Path, sheet: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in workbook.sheetnames:
            raise ValueError(f"{path.name}: falta la hoja {sheet}")
        rows = workbook[sheet].values
        headers = next(rows, None)
        if not headers:
            raise ValueError(f"{path.name}: hoja {sheet} vacía")
        if len(set(headers)) != len(headers) or not REQUIRED_COLUMNS[sheet].issubset(set(headers)):
            missing = sorted(REQUIRED_COLUMNS[sheet] - set(headers))
            raise ValueError(f"{path.name}: columnas duplicadas o ausentes en {sheet}: {missing}")
        return [dict(zip(headers, row)) for row in rows if any(cell is not None for cell in row)]
    finally:
        workbook.close()


def line_key(order_id: Any, line_id: Any) -> str:
    return f"{str(order_id or '').strip().upper()}:{str(line_id or '').strip().removesuffix('.0')}"


def valid_nonnegative_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and value >= 0


def date_only(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return None


def classify_email(subject: str, body: str) -> tuple[str, str | None, str | None]:
    """Return intent, requested date and explicitly superseded message ID."""
    text = fold(f"{subject} {body}")
    supersedes = SUPERSEDES_RE.search(body)
    supersedes_id = supersedes.group(1).lower() if supersedes else None
    no_change = any(phrase in text for phrase in (
        "no solicitamos", "no estamos solicitando", "no solicitamos otro cambio",
        "mantenemos lo acordado", "solo un acuse", "este correo es solo un acuse",
    ))
    if no_change:
        intent = "acuse" if "acuse" in text or "mantenemos" in text or "gracias" in text else "seguimiento"
    elif any(word in text for word in ("cancelar", "cancelacion")):
        intent = "cancelacion"
    elif any(word in text for word in (
        "solicitamos", "solicitud", "cambiad", "cambiar la fecha", "mover la fecha",
        "nueva fecha", "ajustar", "adelantar", "rectifico", "rectificacion",
    )):
        intent = "cambio_fecha"
    elif any(word in text for word in ("estado", "seguimiento", "pendientes", "revisad", "actualizacion")):
        intent = "seguimiento"
    else:
        intent = "informativo"
    requested = None
    if intent == "cambio_fecha":
        match = DATE_RE.search(body)
        if match:
            try:
                requested = date(int(match.group(3)), int(match.group(2)), int(match.group(1))).isoformat()
            except ValueError:
                requested = None
        elif "manana" in text or "hoy" in text:
            requested = "relative:tomorrow" if "manana" in text else "relative:today"
    return intent, requested, supersedes_id


def extract_references(subject: str, body: str, order_lines: dict[str, list[str]]) -> tuple[list[dict[str, str]], list[str]]:
    """Resolve explicit order/line pairs; return unresolved reasons separately."""
    matches = list(ORDER_RE.finditer(body))
    if not matches:
        matches = list(ORDER_RE.finditer(subject))
        source = subject
    else:
        source = body
    links: list[dict[str, str]] = []
    issues: list[str] = []
    for index, match in enumerate(matches):
        order = match.group().upper()
        next_start = matches[index + 1].start() if index + 1 < len(matches) else len(source)
        segment = source[match.end():next_start]
        line_match = LINE_RE.search(segment) or SLASH_LINE_RE.search(segment)
        if not line_match and index == 0:
            line_match = LINE_RE.search(source[max(0, match.start() - 90):match.start()])
        if not line_match and source is subject:
            line_match = LINE_RE.search(body)
        if line_match:
            key = line_key(order, line_match.group(1))
            if key in order_lines.get(order, []):
                links.append({"case_id": f"line:{key}", "match": "pedido+línea"})
            else:
                issues.append(f"{order} / {line_match.group(1)} no existe en el ERP")
        else:
            candidates = order_lines.get(order, [])
            if len(candidates) == 1:
                links.append({"case_id": f"line:{candidates[0]}", "match": "pedido único"})
            elif not candidates:
                issues.append(f"{order} no existe en el ERP")
            else:
                issues.append(f"{order} tiene {len(candidates)} líneas; falta línea")
    if not matches:
        issues.append("Correo sin referencia de pedido")
    deduped = list({link["case_id"]: link for link in links}.values())
    return deduped, issues


def candidate_lines(subject: str, body: str, sender: str, lines: dict[str, dict[str, Any]], customers: dict[str, dict[str, Any]]) -> list[str]:
    """Suggestions only; never silently link a message without order reference."""
    text = f"{subject} {body}"
    skus = {m.group().upper() for m in SKU_RE.finditer(text)}
    if not skus:
        return []
    known_clients = {cid for cid, c in customers.items() if (c.get("email") or "").casefold() == sender.casefold()}
    candidates = [key for key, row in lines.items() if row.get("sku") in skus and (not known_clients or row.get("cliente_id") in known_clients)]
    return sorted(candidates)[:8]


def connect(path: Path) -> sqlite3.Connection:
    path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(path)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA journal_mode=WAL")
    connection.executescript("""
        CREATE TABLE IF NOT EXISTS messages (
          message_id TEXT PRIMARY KEY, payload_hash TEXT NOT NULL, received_at TEXT,
          sender TEXT, recipient TEXT, subject TEXT, body TEXT
        );
        CREATE TABLE IF NOT EXISTS cases (
          case_id TEXT PRIMARY KEY, kind TEXT NOT NULL, order_id TEXT, line_id TEXT,
          client_id TEXT, client_name TEXT, priority TEXT NOT NULL, priority_rank INTEGER NOT NULL,
          due_date TEXT, pending REAL, reason TEXT, action TEXT, attention INTEGER NOT NULL,
          unresolved INTEGER NOT NULL, sort_time TEXT, detail_json TEXT NOT NULL,
          search_text TEXT NOT NULL DEFAULT ''
        );
        CREATE INDEX IF NOT EXISTS cases_queue ON cases(attention, priority_rank, due_date, case_id);
        CREATE INDEX IF NOT EXISTS cases_client ON cases(client_id, priority);
        CREATE TABLE IF NOT EXISTS runs (
          run_id INTEGER PRIMARY KEY AUTOINCREMENT, dataset TEXT NOT NULL, as_of TEXT NOT NULL,
          orders_file TEXT NOT NULL, email_file TEXT NOT NULL, orders_sha256 TEXT NOT NULL,
          added_messages INTEGER NOT NULL, repeated_messages INTEGER NOT NULL,
          conflicting_messages INTEGER NOT NULL, active_messages INTEGER NOT NULL,
          total_cases INTEGER NOT NULL, attention_cases INTEGER NOT NULL, unresolved_cases INTEGER NOT NULL,
          output_sha256 TEXT NOT NULL, elapsed_ms REAL NOT NULL, created_at TEXT NOT NULL
        );
    """)
    if "search_text" not in {row[1] for row in connection.execute("PRAGMA table_info(cases)")}:
        connection.execute("ALTER TABLE cases ADD COLUMN search_text TEXT NOT NULL DEFAULT ''")
    return connection


def _consensus(group: list[dict[str, Any]], field: str) -> Any:
    values = {canonical_json(plain(row.get(field))) for row in group}
    return group[0].get(field) if len(values) == 1 else None


def build_lines(order_rows: list[dict[str, Any]], customer_rows: list[dict[str, Any]]) -> tuple[dict[str, dict[str, Any]], dict[str, dict[str, Any]], dict[str, list[str]]]:
    customers = {str(row.get("cliente_id")): {key: plain(value) for key, value in row.items()} for row in customer_rows}
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in order_rows:
        groups[line_key(row.get("pedido_id"), row.get("linea_id"))].append(row)
    lines: dict[str, dict[str, Any]] = {}
    by_order: dict[str, list[str]] = defaultdict(list)
    for key, group in groups.items():
        order_id, line_id = key.split(":", 1)
        variants = {canonical_json({k: plain(v) for k, v in row.items()}) for row in group}
        fields = {field: _consensus(group, field) for field in group[0]}
        ordered = fields.get("uds_pedidas")
        shipped = fields.get("uds_enviadas")
        issues: list[str] = []
        if len(variants) > 1:
            issues.append("filas contradictorias en el ERP")
        elif len(group) > 1:
            issues.append("duplicado idéntico consolidado")
        if not valid_nonnegative_number(ordered) or not valid_nonnegative_number(shipped):
            issues.append("unidades desconocidas o inválidas")
            pending = None
        elif shipped > ordered:
            issues.append("unidades enviadas mayores que pedidas")
            pending = None
        else:
            pending = ordered - shipped
        commitment = date_only(fields.get("fecha_compromiso"))
        if not commitment:
            issues.append("fecha de compromiso desconocida")
        cid = str(fields.get("cliente_id") or "")
        if cid not in customers:
            issues.append("cliente no encontrado en el maestro")
        line = {
            "case_id": f"line:{key}", "kind": "line", "order_id": order_id, "line_id": line_id,
            "cliente_id": cid, "cliente": customers.get(cid, {}).get("nombre") or "Cliente sin identificar",
            "sku": fields.get("sku"), "color": fields.get("color"), "talla": fields.get("talla"),
            "uds_pedidas": ordered, "uds_enviadas": shipped, "pendientes": pending,
            "precio_unitario_eur": fields.get("precio_unitario_eur"), "fecha_compromiso": commitment,
            "issues": issues, "source_rows": [{k: plain(v) for k, v in row.items()} for row in group],
            "emails": [], "requested_date": None, "request_conflict": False,
        }
        lines[key] = line
        by_order[order_id].append(key)
    return lines, customers, by_order


def _sort_case(case: dict[str, Any]) -> tuple[Any, ...]:
    # Priority and dates are explicit; unknown dates sort after known dates.
    return (case["priority_rank"], case.get("fecha_compromiso") or "9999-12-31", case["case_id"])


def build_cases(order_rows: list[dict[str, Any]], customer_rows: list[dict[str, Any]], messages: list[dict[str, Any]], as_of: datetime, *, include_inactive: bool = False) -> list[dict[str, Any]]:
    lines, customers, by_order = build_lines(order_rows, customer_rows)
    active_messages = sorted((m for m in messages if m["received_at"] <= as_of.isoformat(timespec="minutes")), key=lambda m: (m["received_at"], m["message_id"]))
    unresolved: list[dict[str, Any]] = []
    for msg in active_messages:
        intent, requested, supersedes = classify_email(msg["subject"] or "", msg["body"] or "")
        if requested and requested.startswith("relative:"):
            received_day = date.fromisoformat(msg["received_at"][:10])
            requested = (received_day + timedelta(days=1 if requested.endswith("tomorrow") else 0)).isoformat()
        links, issues = extract_references(msg["subject"] or "", msg["body"] or "", by_order)
        evidence = {"message_id": msg["message_id"], "received_at": msg["received_at"], "from": msg["sender"],
                    "to": msg["recipient"], "subject": msg["subject"], "body": msg["body"], "intent": intent,
                    "requested_date": requested, "supersedes": supersedes, "links": links}
        for link in links:
            key = link["case_id"].removeprefix("line:")
            line = lines[key]
            item = dict(evidence, match=link["match"])
            known_email = customers.get(line["cliente_id"], {}).get("email") or ""
            item["sender_unverified"] = bool(known_email and known_email.casefold() != (msg["sender"] or "").casefold())
            line["emails"].append(item)
        if issues:
            suggestions = candidate_lines(msg["subject"] or "", msg["body"] or "", msg["sender"] or "", lines, customers)
            referenced_order = ORDER_RE.search(f"{msg['subject']} {msg['body']}")
            referenced_line = LINE_RE.search(msg["body"] or "")
            sender_clients = [(cid, customer) for cid, customer in customers.items()
                              if customer.get("email") and customer["email"].casefold() == (msg["sender"] or "").casefold()]
            sender_client_id, sender_client = sender_clients[0] if len(sender_clients) == 1 else (None, None)
            unresolved.append({"case_id": f"email:{msg['message_id']}", "kind": "unresolved_email",
                               "order_id": referenced_order.group().upper() if referenced_order else None,
                               "line_id": referenced_line.group(1) if referenced_line else None,
                               "cliente_id": sender_client_id,
                               "cliente": sender_client["nombre"] if sender_client else "Sin identificar", "sku": None, "color": None,
                               "talla": None, "uds_pedidas": None, "uds_enviadas": None, "pendientes": None,
                               "precio_unitario_eur": None, "fecha_compromiso": None, "requested_date": requested,
                               "issues": issues, "suggestions": suggestions, "emails": [evidence], "request_conflict": False})
    cases: list[dict[str, Any]] = []
    today = as_of.date()
    for line in lines.values():
        emails = line["emails"]
        superseded = {m["supersedes"] for m in emails if m["supersedes"]}
        requests = [m for m in emails if m["intent"] == "cambio_fecha" and m["message_id"] not in superseded]
        if requests:
            line["requested_date"] = requests[-1]["requested_date"]
            line["request_conflict"] = len(requests) > 1
        intents = {m["intent"] for m in emails}
        changed_contact = any("nuevo contacto" in fold(m["body"] or "") or m["sender_unverified"] for m in emails if m["intent"] in ("cambio_fecha", "cancelacion"))
        if changed_contact:
            line["issues"].append("contacto o remitente por verificar")
        if line["request_conflict"]:
            line["issues"].append("varias peticiones sin rectificación explícita")
        has_anomaly = any(issue not in ("duplicado idéntico consolidado",) for issue in line["issues"])
        pending = line["pendientes"]
        commitment = line["fecha_compromiso"]
        overdue = bool(pending and commitment and date.fromisoformat(commitment) < today)
        due_today = bool(pending and commitment and date.fromisoformat(commitment) == today)
        due_soon = bool(pending and commitment and 0 < (date.fromisoformat(commitment) - today).days <= 2)
        has_request = bool(requests)
        cancellation = "cancelacion" in intents
        followup = "seguimiento" in intents
        urgent_request = bool(line["requested_date"] and date.fromisoformat(line["requested_date"]) <= today)
        if cancellation:
            priority, rank = "Alta", 0
        elif urgent_request:
            priority, rank = "Alta", 1
        elif changed_contact:
            priority, rank = "Alta", 2
        elif has_anomaly and pending is None:
            priority, rank = "Alta", 3
        elif overdue:
            priority, rank = "Alta", 4
        elif due_today:
            priority, rank = "Alta", 5
        elif has_request:
            priority, rank = "Media", 10
        elif has_anomaly:
            priority, rank = "Media", 11
        elif due_soon:
            priority, rank = "Media", 12
        elif followup:
            priority, rank = "Media", 13
        else:
            priority, rank = "Baja", 20
        attention = bool(pending is None or pending > 0 or cancellation or has_request or has_anomaly or followup)
        if cancellation:
            reason, action = "Solicitud de cancelación", "Comprobar expedición y validar la cancelación antes de cambiar el ERP"
        elif urgent_request:
            reason, action = "Adelanto solicitado para hoy", "Comprobar stock y logística antes de confirmar una fecha"
        elif changed_contact:
            reason, action = "Contacto por verificar", "Verificar autorización por un canal conocido antes de responder o cambiar el ERP"
        elif pending is None:
            reason, action = "Datos del ERP contradictorios", "Revisar la línea en Navision antes de calcular o prometer unidades"
        elif overdue:
            reason, action = "Compromiso vencido", "Verificar expedición y comunicar un estado contrastado"
        elif due_today:
            reason, action = "Compromiso para hoy", "Comprobar salida de mercancía"
        elif has_request:
            reason, action = "Cambio de fecha solicitado", "Comprobar viabilidad, confirmar y registrar el cambio aprobado en el ERP"
        elif has_anomaly:
            reason, action = "Dato incompleto", "Revisar el origen antes de tomar una decisión"
        elif due_soon:
            reason, action = "Compromiso próximo", "Preparar la expedición y confirmar disponibilidad"
        elif followup:
            reason, action = "Consulta de estado", "Comprobar expedición y responder con unidades pendientes verificadas"
        else:
            reason, action = "Unidades pendientes", "Planificar la siguiente expedición"
        if not attention:
            reason, action = "Línea servida sin incidencia", "Sin acción operativa pendiente"
        line.update({"priority": priority, "priority_rank": rank, "attention": attention,
                     "unresolved": False, "reason": reason, "action": action, "overdue": overdue,
                     "due_today": due_today, "request_count": len(requests), "sort_time": emails[-1]["received_at"] if emails else None})
        line["label"] = case_label(line)
        if attention or include_inactive:
            cases.append(line)
    for item in unresolved:
        item.update({"priority": "Alta" if item["emails"][0]["intent"] in ("cambio_fecha", "cancelacion") else "Media",
                     "priority_rank": 2 if item["emails"][0]["intent"] in ("cambio_fecha", "cancelacion") else 13,
                     "attention": True, "unresolved": True, "reason": item["issues"][0],
                     "action": "Investigar la referencia y asociar solo tras verificación humana",
                     "sort_time": item["emails"][0]["received_at"]})
        item["label"] = case_label(item)
        cases.append(item)
    return sorted(cases, key=_sort_case)


@dataclass(frozen=True)
class RunResult:
    run_id: int
    dataset: str
    as_of: str
    added_messages: int
    repeated_messages: int
    conflicting_messages: int
    active_messages: int
    total_cases: int
    attention_cases: int
    unresolved_cases: int
    output_sha256: str
    elapsed_ms: float


def run_import(db_path: Path, dataset: str, orders_path: Path, email_path: Path, as_of: datetime) -> RunResult:
    started = time.perf_counter()
    order_rows = load_sheet(orders_path, "Pedidos")
    customer_rows = load_sheet(orders_path, "Clientes")
    email_rows = load_sheet(email_path, "Correos")
    orders_sha = digest({"orders": order_rows, "customers": customer_rows})
    connection = connect(db_path)
    try:
        prior_dataset = connection.execute("SELECT dataset FROM runs ORDER BY run_id DESC LIMIT 1").fetchone()
        if prior_dataset and prior_dataset["dataset"] != dataset:
            raise ValueError(f"La base {db_path} pertenece a {prior_dataset['dataset']}, no a {dataset}")
        existing = {row["message_id"]: row["payload_hash"] for row in connection.execute("SELECT message_id,payload_hash FROM messages")}
        added = repeated = conflicting = 0
        with connection:
            for row in email_rows:
                mid = str(row.get("message_id") or "").strip().lower()
                if not mid:
                    conflicting += 1
                    continue
                received = row.get("received_at")
                if not isinstance(received, datetime):
                    conflicting += 1
                    continue
                payload = {"received_at": plain(received), "sender": row.get("from") or "", "recipient": row.get("to") or "",
                           "subject": row.get("subject") or "", "body": row.get("body") or ""}
                fingerprint = digest(payload)
                if mid in existing:
                    if existing[mid] == fingerprint:
                        repeated += 1
                    else:
                        conflicting += 1
                    continue
                connection.execute("INSERT INTO messages VALUES (?,?,?,?,?,?,?)", (mid, fingerprint, payload["received_at"],
                                   payload["sender"], payload["recipient"], payload["subject"], payload["body"]))
                existing[mid] = fingerprint
                added += 1
            messages = [dict(row) for row in connection.execute("SELECT message_id,received_at,sender,recipient,subject,body FROM messages")]
            all_cases = build_cases(order_rows, customer_rows, messages, as_of, include_inactive=True)
            cases = [case for case in all_cases if case["attention"]]
            customer_emails = {str(row.get("cliente_id")): row.get("email") or "" for row in customer_rows}
            connection.execute("DELETE FROM cases")
            connection.executemany("""INSERT INTO cases
              (case_id,kind,order_id,line_id,client_id,client_name,priority,priority_rank,due_date,
               pending,reason,action,attention,unresolved,sort_time,detail_json,search_text)
              VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", [
                (c["case_id"], c["kind"], c.get("order_id"), c.get("line_id"), c.get("cliente_id"), c.get("cliente"),
                 c["priority"], c["priority_rank"], c.get("fecha_compromiso"), c.get("pendientes"), c["reason"],
                 c["action"], int(c["attention"]), int(c["unresolved"]), c.get("sort_time"), canonical_json(c),
                 case_search_text(c, customer_emails.get(c.get("cliente_id"), ""))) for c in all_cases
            ])
            output_sha = digest(cases)
            active_count = sum(m["received_at"] <= as_of.isoformat(timespec="minutes") for m in messages)
            elapsed = round((time.perf_counter() - started) * 1000, 2)
            cursor = connection.execute("""INSERT INTO runs(dataset,as_of,orders_file,email_file,orders_sha256,
              added_messages,repeated_messages,conflicting_messages,active_messages,total_cases,attention_cases,
              unresolved_cases,output_sha256,elapsed_ms,created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
              (dataset, as_of.isoformat(timespec="minutes"), orders_path.name, email_path.name, orders_sha, added,
               repeated, conflicting, active_count, len(cases), len(cases), sum(c["unresolved"] for c in cases),
               output_sha, elapsed, datetime.now().isoformat(timespec="seconds")))
            run_id = cursor.lastrowid
            connection.execute("PRAGMA user_version=3")
        return RunResult(run_id, dataset, as_of.isoformat(timespec="minutes"), added, repeated, conflicting,
                         active_count, len(cases), len(cases), sum(c["unresolved"] for c in cases), output_sha, elapsed)
    finally:
        connection.close()


def summary(db_path: Path) -> dict[str, Any]:
    connection = connect(db_path)
    try:
        latest = connection.execute("SELECT * FROM runs ORDER BY run_id DESC LIMIT 1").fetchone()
        counts = {row["priority"]: row["n"] for row in connection.execute("SELECT priority,COUNT(*) n FROM cases WHERE attention=1 GROUP BY priority")}
        clients = [dict(row) for row in connection.execute("SELECT DISTINCT client_id id,client_name name FROM cases WHERE client_id IS NOT NULL ORDER BY name")]
        records = connection.execute("SELECT COUNT(*) FROM cases").fetchone()[0]
        return {"latest_run": dict(latest) if latest else None, "priorities": counts, "clients": clients,
                "total_records": records}
    finally:
        connection.close()


def list_cases(db_path: Path, *, client: str | None = None, priority: str | None = None, view: str = "all", search: str = "", page: int = 1, page_size: int = 6) -> dict[str, Any]:
    connection = connect(db_path)
    try:
        clauses = [] if view == "records" else ["attention=1"]
        params: list[Any] = []
        if client:
            clauses.append("client_id=?")
            params.append(client)
        if priority:
            # La prioridad solo tiene sentido para casos activos; «Servido» es un estado.
            clauses.append("attention=1")
            clauses.append("priority=?")
            params.append(priority)
        if view == "unresolved":
            clauses.append("unresolved=1")
        elif view == "high":
            clauses.append("priority='Alta'")
        if search.strip():
            terms = fold(search.strip()).split()
            for term in terms:
                if term in ("alta", "media", "baja"):
                    clauses.extend(("attention=1", "priority=?"))
                    params.append(term.capitalize())
                elif term in ("servido", "servida"):
                    clauses.append("attention=0")
                else:
                    clauses.append("search_text LIKE ? ESCAPE '\\'")
                    escaped = term.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
                    params.append(f"%{escaped}%")
        where = " AND ".join(clauses) if clauses else "1=1"
        total = connection.execute(f"SELECT COUNT(*) FROM cases WHERE {where}", params).fetchone()[0]
        rows = connection.execute(f"""SELECT case_id,kind,order_id,line_id,client_id,client_name,priority,
          due_date,pending,reason,action,attention,unresolved FROM cases WHERE {where}
          ORDER BY priority_rank,due_date IS NULL,due_date,sort_time DESC,case_id LIMIT ? OFFSET ?""",
          [*params, page_size, (max(1, page) - 1) * page_size]).fetchall()
        return {"total": total, "page": max(1, page), "page_size": page_size, "items": [dict(row) for row in rows]}
    finally:
        connection.close()


def get_case(db_path: Path, case_id: str) -> dict[str, Any] | None:
    connection = connect(db_path)
    try:
        row = connection.execute("SELECT detail_json FROM cases WHERE case_id=?", (case_id,)).fetchone()
        return json.loads(row[0]) if row else None
    finally:
        connection.close()


def recent_runs(db_path: Path, limit: int = 20) -> list[dict[str, Any]]:
    connection = connect(db_path)
    try:
        return [dict(row) for row in connection.execute("SELECT * FROM runs ORDER BY run_id DESC LIMIT ?", (limit,))]
    finally:
        connection.close()
